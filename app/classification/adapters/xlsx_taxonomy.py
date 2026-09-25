from pathlib import Path

from openpyxl import load_workbook

from app.classification.domain.models import Category, Taxonomy


HEADERS = ("uidentifier", "family", "name", "parent", "is_generic_or_other")


def read_categories(path: Path) -> list[Category]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = tuple(next(rows, ()))
        if headers != HEADERS:
            raise ValueError(f"Columnas de taxonomía inválidas: {headers!r}")
        categories = [_category(row, row_number) for row_number, row in enumerate(rows, start=2)]
    finally:
        workbook.close()
    identifiers = [category.uidentifier for category in categories]
    if len(identifiers) != len(set(identifiers)):
        raise ValueError("La taxonomía contiene uidentifier duplicados")
    return categories


def read_taxonomy(path: Path) -> Taxonomy:
    return Taxonomy(tuple(read_categories(path)))


def _category(row: tuple, row_number: int) -> Category:
    if len(row) != len(HEADERS) or any(not isinstance(value, str) or not value.strip() for value in row[:4]):
        raise ValueError(f"Fila de taxonomía inválida: {row_number}")
    generic = row[4]
    if isinstance(generic, str):
        if generic.casefold() not in ("true", "false"):
            raise ValueError(f"Valor is_generic_or_other inválido en fila {row_number}")
        generic = generic.casefold() == "true"
    elif not isinstance(generic, bool):
        raise ValueError(f"Valor is_generic_or_other inválido en fila {row_number}")
    return Category(*(value.strip() for value in row[:4]), generic)
